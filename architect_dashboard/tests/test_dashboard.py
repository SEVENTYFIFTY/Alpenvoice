import io
from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from architect_dashboard import config, db, excel_io, gdrive, progress


@pytest.fixture(autouse=True)
def temp_db(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DB_PATH", str(tmp_path / "test.db"))
    monkeypatch.setattr(config, "GOOGLE_SERVICE_ACCOUNT_FILE", "")
    db.init_db()


@pytest.fixture
def client():
    from architect_dashboard.app import app
    with TestClient(app) as c:
        yield c


def _project(conn, template="sia112", **extra):
    return db.create_project(conn, {"code": "25-001", "name": "Test House", **extra}, template=template)


# ---------------------------------------------------------------------------
# progress maths
# ---------------------------------------------------------------------------

def test_overall_progress_is_weighted_by_phase():
    phases = [{"progress": 100, "weight": 1}, {"progress": 0, "weight": 3}]
    assert progress.weighted_progress(phases) == 25


def test_expected_progress_is_linear_between_dates():
    today = date(2026, 1, 11)
    assert progress.expected_progress("2026-01-01", "2026-01-21", today) == 50
    assert progress.expected_progress("2026-01-01", None, today) is None


def test_phase_health():
    today = date(2026, 1, 11)
    base = {"planned_start": "2026-01-01", "planned_end": "2026-01-21"}
    assert progress.phase_health({**base, "progress": 50}, today) == "on_track"
    assert progress.phase_health({**base, "progress": 20}, today) == "at_risk"
    assert progress.phase_health({**base, "progress": 100}, today) == "done"
    assert progress.phase_health({**base, "planned_end": "2026-01-10", "progress": 90}, today) == "overdue"


def test_template_creates_phases_and_future_phases_do_not_make_project_not_started():
    with db.connect() as conn:
        project_id = _project(conn)
        phases = db.list_phases(conn, project_id)
        assert [p["code"] for p in phases][:3] == ["21", "31", "32"]
        db.set_phase_progress(conn, phases[0]["id"], 40)
        summary = progress.project_summary(conn, db.get_project(conn, project_id), date.today())
    assert summary["health"] == "on_track"
    assert summary["current_phase"] == "Preliminary studies"


def test_daily_delta_uses_progress_log():
    today = date.today()
    yesterday = (today - timedelta(days=1)).isoformat() + "T10:00:00"
    with db.connect() as conn:
        project_id = _project(conn, template=None)
        phase_id = db.upsert_phase(conn, project_id, {"name": "Design", "weight": 1})
        db.set_phase_progress(conn, phase_id, 30, logged_at=yesterday)
        db.set_phase_progress(conn, phase_id, 45)
        summary = progress.project_summary(conn, db.get_project(conn, project_id), today)
    assert summary["progress"] == 45
    assert summary["delta_today"] == 15
    assert summary["history"][-2]["progress"] == 30
    assert summary["history"][-1]["progress"] == 45


def test_progress_is_clamped_and_completes_phase():
    with db.connect() as conn:
        project_id = _project(conn, template=None)
        phase_id = db.upsert_phase(conn, project_id, {"name": "Permit"})
        db.set_phase_progress(conn, phase_id, 140)
        phase = db.get_phase(conn, phase_id)
    assert phase["progress"] == 100 and phase["status"] == "done"


def test_drawing_stage_sets_default_progress():
    with db.connect() as conn:
        project_id = _project(conn, template=None)
        drawing_id = db.upsert_drawing(conn, project_id, {"number": "A-101", "stage": "review"})
        assert db.get_drawing(conn, drawing_id)["progress"] == 75
        db.upsert_drawing(conn, project_id, {"number": "A-101", "stage": "issued"})
        assert db.get_drawing(conn, drawing_id)["progress"] == 100
        assert len(db.list_drawings(conn, project_id)) == 1


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def _workbook() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Projects")
    ws.append(["Code", "Name", "Lead Architect", "Start Date", "Due Date", "Template"])
    ws.append(["26-001", "Library", "Eva Muster", date(2026, 1, 1), "31.12.2026", "SIA112"])
    ws = wb.create_sheet("Phases")
    ws.append(["Project", "Phase", "Progress %", "Responsible"])
    ws.append(["26-001", "Preliminary design", 0.4, "Tom Beispiel"])
    ws.append(["26-001", "Interior design", 20, None])
    ws.append(["99-999", "Ghost phase", 10, None])
    ws = wb.create_sheet("Drawings")
    ws.append(["Project", "Drawing No", "Title", "Stage", "Phase"])
    ws.append(["26-001", "A-100", "Site plan", "In Progress", "31"])
    ws.append(["26-001", "A-101", "Plan", "finished", None])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_excel_import_upserts_and_reports_errors():
    with db.connect() as conn:
        report = excel_io.import_workbook(conn, _workbook())
        project_id = db.project_id_by_code(conn, "26-001")
        project = db.get_project(conn, project_id)
        phases = {p["name"]: p for p in db.list_phases(conn, project_id)}
        drawings = db.list_drawings(conn, project_id)

    assert report["projects"] == 1 and report["phases"] == 2 and report["drawings"] == 1
    assert len(report["errors"]) == 2  # unknown project + unknown stage
    assert project["lead_name"] == "Eva Muster"
    assert project["due_date"] == "2026-12-31"
    assert phases["Preliminary design"]["progress"] == 40
    assert phases["Preliminary design"]["assignee_name"] == "Tom Beispiel"
    assert phases["Interior design"]["progress"] == 20  # new phase appended to the template
    assert len(phases) == 9
    assert drawings[0]["phase_name"] == "Preliminary design"

    # importing again doesn't duplicate anything
    with db.connect() as conn:
        excel_io.import_workbook(conn, _workbook())
        assert len(db.list_phases(conn, project_id)) == 9
        assert len(db.list_projects(conn)) == 1


def test_excel_export_round_trips():
    with db.connect() as conn:
        excel_io.import_workbook(conn, _workbook())
        content = excel_io.export_workbook(conn)
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Summary", "Team", "Projects", "Phases", "Drawings"]
    assert wb["Summary"]["A2"].value == "26-001"

    # a second database built from the export ends up with the same progress
    with db.connect() as conn:
        conn.executescript("DELETE FROM projects; DELETE FROM members;")
        excel_io.import_workbook(conn, content)
        project_id = db.project_id_by_code(conn, "26-001")
        phases = {p["name"]: p["progress"] for p in db.list_phases(conn, project_id)}
    assert phases["Preliminary design"] == 40


# ---------------------------------------------------------------------------
# Google Drive
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("link", [
    "https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOp/edit#gid=0",
    "https://drive.google.com/drive/folders/1AbCdEfGhIjKlMnOp?usp=sharing",
    "https://drive.google.com/open?id=1AbCdEfGhIjKlMnOp",
    "1AbCdEfGhIjKlMnOp",
])
def test_extract_drive_id(link):
    assert gdrive.extract_id(link) == "1AbCdEfGhIjKlMnOp"


def test_extract_drive_id_rejects_garbage():
    with pytest.raises(gdrive.DriveError):
        gdrive.extract_id("not a link")


class FakeResponse:
    def __init__(self, status=200, json=None, content=b""):
        self.status_code, self._json, self.content, self.text = status, json, content, ""

    def json(self):
        return self._json


def test_drive_sync_exports_google_sheet_and_imports(monkeypatch):
    calls = []

    class FakeSession:
        def get(self, url, params=None, timeout=None):
            calls.append((url, params))
            if url.endswith("/export"):
                return FakeResponse(content=_workbook())
            return FakeResponse(json={"id": "sheet1234567", "name": "Office tracker", "mimeType": gdrive.SHEET_MIME})

    monkeypatch.setattr(gdrive, "session", lambda: FakeSession())
    with db.connect() as conn:
        source_id = db.add_sync_source(conn, "sheet1234567")
    result = gdrive.sync_source({"id": source_id, "file_id": "sheet1234567", "name": None})

    assert result.startswith("OK: 1 projects")
    assert calls[1][1]["mimeType"] == gdrive.XLSX_MIME
    with db.connect() as conn:
        source = db.list_sync_sources(conn)[0]
        assert source["name"] == "Office tracker"
        assert db.project_id_by_code(conn, "26-001")


def test_drive_sync_records_errors(monkeypatch):
    class FakeSession:
        def get(self, url, params=None, timeout=None):
            return FakeResponse(status=404, json={"error": {"message": "File not found"}})

    monkeypatch.setattr(gdrive, "session", lambda: FakeSession())
    with db.connect() as conn:
        source_id = db.add_sync_source(conn, "missing12345")
    result = gdrive.sync_source({"id": source_id, "file_id": "missing12345", "name": None})
    assert "File not found" in result
    with db.connect() as conn:
        assert db.list_sync_sources(conn)[0]["last_result"].startswith("Error")


# ---------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------

def test_api_project_lifecycle(client):
    member = client.post("/api/members", json={"name": "Anna", "role": "Principal"}).json()
    res = client.post("/api/projects", json={
        "code": "26-010", "name": "Museum", "lead_id": member["id"], "template": "international",
        "start_date": "2026-01-01", "due_date": "2027-01-01",
        "drive_folder_id": "https://drive.google.com/drive/folders/1AbCdEfGhIjKlMnOp",
    })
    assert res.status_code == 201
    project_id = res.json()["id"]
    assert client.post("/api/projects", json={"code": "26-010", "name": "Dup"}).status_code == 409

    project = client.get(f"/api/projects/{project_id}").json()
    assert project["drive_folder_id"] == "1AbCdEfGhIjKlMnOp"
    phase = project["phases"][0]
    assert client.patch(f"/api/phases/{phase['id']}", json={"progress": 150}).status_code == 422
    ok = client.patch(f"/api/phases/{phase['id']}",
                      json={"progress": 60, "note": "Massing approved", "member_id": member["id"]})
    assert ok.status_code == 200

    drawing = client.post(f"/api/projects/{project_id}/drawings", json={"number": "A-1", "stage": "draft"})
    assert drawing.status_code == 201
    assert client.patch(f"/api/drawings/{drawing.json()['id']}", json={"stage": "bogus"}).status_code == 422

    dash = client.get("/api/dashboard").json()
    assert dash["kpis"]["active_projects"] == 1
    assert dash["projects"][0]["progress"] == 6.0  # 60% of a 10% weight phase
    update = next(a for a in dash["activity"] if a["note"] == "Massing approved")
    assert update["member_name"] == "Anna"

    files = client.get(f"/api/projects/{project_id}/drive-files").json()
    assert files["files"] == [] and "not configured" in files["message"]

    assert client.delete(f"/api/projects/{project_id}").status_code == 200
    assert client.get(f"/api/projects/{project_id}").status_code == 404


def test_api_excel_endpoints(client):
    res = client.post("/api/import/excel", files={"file": ("tracker.xlsx", _workbook(), gdrive.XLSX_MIME)})
    assert res.status_code == 200 and res.json()["projects"] == 1
    assert client.post("/api/import/excel", files={"file": ("notes.txt", b"hi", "text/plain")}).status_code == 422
    template = client.get("/api/export/template")
    assert load_workbook(io.BytesIO(template.content)).sheetnames[:4] == ["Team", "Projects", "Phases", "Drawings"]
    assert client.get("/api/export/excel").status_code == 200


def test_index_served(client):
    res = client.get("/")
    assert res.status_code == 200 and "Studio Dashboard" in res.text
