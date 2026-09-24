"""Excel import/export.

One workbook format is used everywhere: the downloadable template, the export,
uploads, and spreadsheets pulled from Google Drive. Sheets (all optional):

  Team        Name | Role | Email | Status
  Projects    Code | Name | Client | Location | Lead | Status | Start | Due | Drive Folder | Template | Blocker | Notes
  Phases      Project | Code | Phase | Weight | Planned Start | Planned End | Progress | Status | Assignee
  Milestones  Project | Phase | Milestone | Due | Done
  Drawings    Project | Number | Title | Phase | Discipline | Scale | Revision | Stage | Progress | Assignee | Due
  Contacts    Project | Name | Type | Role | Email | Phone | Notify

Rows are matched on their natural keys (team member name, project code,
project + phase name, project + drawing number), so re-importing the same
workbook updates records instead of duplicating them.
"""
import io
from datetime import date, datetime
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db, progress
from .phases import DRAWING_STAGES, TEMPLATES

SHEETS = {
    "Team": ["Name", "Role", "Email", "Status"],
    "Projects": ["Code", "Name", "Client", "Location", "Lead", "Status", "Start", "Due",
                 "Drive Folder", "Template", "Blocker", "Notes"],
    "Phases": ["Project", "Code", "Phase", "Weight", "Planned Start", "Planned End",
               "Progress", "Status", "Assignee"],
    "Milestones": ["Project", "Phase", "Milestone", "Due", "Done"],
    "Drawings": ["Project", "Number", "Title", "Phase", "Discipline", "Scale", "Revision",
                 "Stage", "Progress", "Assignee", "Due"],
    "Contacts": ["Project", "Name", "Type", "Role", "Email", "Phone", "Notify"],
}

TRUE_WORDS = {"yes", "y", "true", "1", "x", "✓", "done", "ja", "oui", "sim", "si"}

# Accepted alternative header spellings -> canonical header
ALIASES = {
    "project code": "project", "project name": "name", "phase name": "phase",
    "lead architect": "lead", "start date": "start", "due date": "due", "deadline": "due",
    "drive folder id": "drive folder", "drive": "drive folder", "%": "progress",
    "progress %": "progress", "percent": "progress", "drawing": "number", "drawing no": "number",
    "drawing number": "number", "rev": "revision", "responsible": "assignee", "owner": "assignee",
    "start planned": "planned start", "end planned": "planned end",
    "task": "milestone", "kind": "type", "notify on phase change": "notify",
    "waiting on": "blocker", "blocked by": "blocker",
}


class ImportReport(dict):
    def __init__(self):
        super().__init__(team=0, projects=0, phases=0, milestones=0, drawings=0, contacts=0, errors=[])

    def summary(self) -> str:
        text = (f"{self['projects']} projects, {self['phases']} phases, {self['milestones']} milestones, "
                f"{self['drawings']} drawings, {self['contacts']} contacts, {self['team']} team members")
        if self["errors"]:
            text += f" — {len(self['errors'])} row(s) skipped"
        return text


def _key(header) -> str:
    text = str(header or "").strip().lower()
    return ALIASES.get(text, text)


def _date(value) -> Optional[str]:
    if value in (None, ""):
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    raise ValueError(f"unrecognised date '{text}'")


def _percent(value) -> Optional[float]:
    if value in (None, ""):
        return None
    if isinstance(value, str):
        value = value.strip().rstrip("%").replace(",", ".")
    number = float(value)
    # Excel percentage-formatted cells arrive as fractions (0.45 for 45%, 1 for 100%)
    if 0 < number <= 1:
        number *= 100
    return number


def _text(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _bool(value) -> Optional[bool]:
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in TRUE_WORDS


def _slug(value) -> Optional[str]:
    text = _text(value)
    return text.lower().replace(" ", "_").replace("-", "_") if text else None


def _sheet_rows(ws) -> list[tuple[int, dict]]:
    iterator = ws.iter_rows(values_only=True)
    try:
        headers = [_key(h) for h in next(iterator)]
    except StopIteration:
        return []
    result = []
    for number, values in enumerate(iterator, start=2):
        if all(v in (None, "") for v in values):
            continue
        result.append((number, dict(zip(headers, values))))
    return result


def _find_sheet(wb, name: str):
    for ws in wb.worksheets:
        if ws.title.strip().lower() == name.lower():
            return ws
    return None


def import_workbook(conn, data: bytes, source: str = "excel") -> ImportReport:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    report = ImportReport()

    def run(sheet: str, handler):
        ws = _find_sheet(wb, sheet)
        if ws is None:
            return
        for number, row in _sheet_rows(ws):
            try:
                handler(row)
                report[sheet.lower()] += 1
            except (ValueError, KeyError, TypeError) as exc:
                report["errors"].append(f"{sheet} row {number}: {exc}")

    def team(row):
        name = _text(row.get("name"))
        if not name:
            raise ValueError("missing Name")
        member_id = db.upsert_member(conn, name, _text(row.get("role")), _text(row.get("email")))
        if _text(row.get("status")):
            db.update_member(conn, member_id, {"status": _text(row.get("status"))})

    def project(row):
        code, name = _text(row.get("code")), _text(row.get("name"))
        if not code:
            raise ValueError("missing Code")
        data = {
            "code": code,
            "name": name,
            "client": _text(row.get("client")),
            "location": _text(row.get("location")),
            "lead_id": db.member_id_by_name(conn, _text(row.get("lead"))),
            "status": _slug(row.get("status")),
            "start_date": _date(row.get("start")),
            "due_date": _date(row.get("due")),
            "drive_folder_id": _text(row.get("drive folder")),
            "notes": _text(row.get("notes")),
            "blocker": _text(row.get("blocker")),
        }
        existing = db.project_id_by_code(conn, code)
        if existing:
            db.update_project(conn, existing, {k: v for k, v in data.items() if v is not None})
        else:
            if not name:
                raise ValueError("missing Name for new project")
            template = _slug(row.get("template")) or db.DEFAULT
            if template == "none":
                template = None
            elif template != db.DEFAULT and template not in TEMPLATES:
                raise ValueError(f"unknown template '{template}'")
            blocker = data.pop("blocker")
            project_id = db.create_project(conn, data, template=template)
            if blocker:
                db.update_project(conn, project_id, {"blocker": blocker})

    def project_for(row) -> int:
        code = _text(row.get("project"))
        project_id = db.project_id_by_code(conn, code) if code else None
        if not project_id:
            raise ValueError(f"unknown project '{code}'")
        return project_id

    def phase(row):
        project_id = project_for(row)
        name = _text(row.get("phase"))
        if not name:
            raise ValueError("missing Phase")
        weight = row.get("weight")
        db.upsert_phase(conn, project_id, {
            "name": name,
            "code": _text(row.get("code")),
            "weight": float(weight) if weight not in (None, "") else None,
            "planned_start": _date(row.get("planned start")),
            "planned_end": _date(row.get("planned end")),
            "status": _slug(row.get("status")),
            "assignee_id": db.member_id_by_name(conn, _text(row.get("assignee"))),
            "progress": _percent(row.get("progress")),
        }, source=source)

    def find_phase(project_id, name):
        if not name:
            return None
        found = conn.execute(
            "SELECT id FROM phases WHERE project_id = ? AND (name = ? OR code = ?)",
            (project_id, name, name),
        ).fetchone()
        return found["id"] if found else None

    def milestone(row):
        project_id = project_for(row)
        phase_id = find_phase(project_id, _text(row.get("phase")))
        if not phase_id:
            raise ValueError(f"unknown phase '{_text(row.get('phase'))}'")
        title = _text(row.get("milestone"))
        if not title:
            raise ValueError("missing Milestone")
        db.upsert_milestone(conn, phase_id, title, _date(row.get("due")), _bool(row.get("done")), source=source)

    def contact(row):
        project_id = project_for(row)
        name = _text(row.get("name"))
        if not name:
            raise ValueError("missing Name")
        kind = _slug(row.get("type"))
        if kind and kind not in db.CONTACT_KINDS:
            raise ValueError(f"unknown type '{kind}' (use: {', '.join(db.CONTACT_KINDS)})")
        db.upsert_contact(conn, project_id, {
            "name": name, "kind": kind, "role": _text(row.get("role")), "email": _text(row.get("email")),
            "phone": _text(row.get("phone")), "notify": _bool(row.get("notify")),
        })

    def drawing(row):
        project_id = project_for(row)
        number = _text(row.get("number"))
        if not number:
            raise ValueError("missing Number")
        stage = _slug(row.get("stage"))
        if stage and stage not in DRAWING_STAGES:
            raise ValueError(f"unknown stage '{stage}' (use: {', '.join(DRAWING_STAGES)})")
        phase_id = find_phase(project_id, _text(row.get("phase")))
        db.upsert_drawing(conn, project_id, {
            "number": number,
            "title": _text(row.get("title")),
            "phase_id": phase_id,
            "discipline": _text(row.get("discipline")),
            "scale": _text(row.get("scale")),
            "revision": _text(row.get("revision")),
            "stage": stage,
            "progress": _percent(row.get("progress")),
            "assignee_id": db.member_id_by_name(conn, _text(row.get("assignee"))),
            "due_date": _date(row.get("due")),
        }, source=source)

    # Order matters: members and projects must exist before phases and drawings.
    run("Team", team)
    run("Projects", project)
    run("Phases", phase)
    run("Milestones", milestone)
    run("Drawings", drawing)
    run("Contacts", contact)
    return report


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="256ABF")


def _write_sheet(ws, headers: list[str], data: list[list]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in data:
        ws.append(row)
    ws.freeze_panes = "A2"
    for index, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[index - 1] or "")) for r in data]) + 2
        ws.column_dimensions[get_column_letter(index)].width = min(width, 45)


def export_workbook(conn, empty: bool = False) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    if empty:
        for name, headers in SHEETS.items():
            _write_sheet(wb.create_sheet(name), headers, [])
        _write_help(wb.create_sheet("Help"))
    else:
        summary = progress.dashboard(conn)
        _write_sheet(wb.create_sheet("Summary"),
                     ["Code", "Project", "Lead", "Status", "Health", "Progress %", "Expected %",
                      "Change today", "Change 7 days", "Current phase", "Due", "Days left",
                      "Drawings", "Issued"],
                     [[p["code"], p["name"], p["lead_name"], p["status"], p["health"], p["progress"],
                       p["expected"], p["delta_today"], p["delta_week"], p["current_phase"],
                       p["due_date"], p["days_left"], p["drawings"]["total"],
                       p["drawings"]["by_stage"]["issued"]] for p in summary["projects"]])

        _write_sheet(wb.create_sheet("Team"), SHEETS["Team"],
                     [[m["name"], m["role"], m["email"], m["status"]] for m in db.list_members(conn)])

        projects = db.list_projects(conn, include_archived=True)
        _write_sheet(wb.create_sheet("Projects"), SHEETS["Projects"],
                     [[p["code"], p["name"], p["client"], p["location"], p["lead_name"], p["status"],
                       p["start_date"], p["due_date"], p["drive_folder_id"], None, p["blocker"], p["notes"]]
                      for p in projects])

        phase_rows, milestone_rows, drawing_rows, contact_rows = [], [], [], []
        for p in projects:
            phase_names = {}
            for ph in db.list_phases(conn, p["id"]):
                phase_names[ph["id"]] = ph["name"]
                phase_rows.append([p["code"], ph["code"], ph["name"], ph["weight"], ph["planned_start"],
                                   ph["planned_end"], ph["progress"], ph["status"], ph["assignee_name"]])
            for m in db.list_milestones(conn, p["id"]):
                milestone_rows.append([p["code"], phase_names[m["phase_id"]], m["title"], m["due_date"],
                                       "yes" if m["done"] else "no"])
            for c in db.list_contacts(conn, p["id"]):
                contact_rows.append([p["code"], c["name"], c["kind"], c["role"], c["email"], c["phone"],
                                     "yes" if c["notify"] else "no"])
            for d in db.list_drawings(conn, p["id"]):
                drawing_rows.append([p["code"], d["number"], d["title"], d["phase_name"], d["discipline"],
                                     d["scale"], d["revision"], d["stage"], d["progress"],
                                     d["assignee_name"], d["due_date"]])
        _write_sheet(wb.create_sheet("Phases"), SHEETS["Phases"], phase_rows)
        _write_sheet(wb.create_sheet("Milestones"), SHEETS["Milestones"], milestone_rows)
        _write_sheet(wb.create_sheet("Drawings"), SHEETS["Drawings"], drawing_rows)
        _write_sheet(wb.create_sheet("Contacts"), SHEETS["Contacts"], contact_rows)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _write_help(ws) -> None:
    lines = [
        ["How to fill this workbook"],
        ["Team: one row per person. Names are how Lead / Assignee columns are matched."],
        ["Projects: Code is the unique key. Template (new projects only): " + ", ".join(TEMPLATES)
         + " or none. Left empty, new projects get the office's default phases."],
        ["Phases: Project = project Code. Progress is a % cell or a number from 0–100 (values up to 1 are read as fractions: 0.5 = 50%). Dates as YYYY-MM-DD or DD.MM.YYYY."],
        ["Milestones: checklist items per phase (Project, Phase name or code, Milestone, Due, Done yes/no). "
         "When a phase has milestones, its progress = share of milestones done."],
        ["Drawings: Stage is one of " + ", ".join(DRAWING_STAGES) + ". Progress defaults from the stage."],
        ["Contacts: Type is one of " + ", ".join(db.CONTACT_KINDS) + ". Notify = yes to get an email draft "
         "when the project moves to a new phase."],
        ["Re-importing updates existing rows; nothing is deleted by an import."],
    ]
    for line in lines:
        ws.append(line)
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 120
